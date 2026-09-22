from __future__ import annotations

import io
from pathlib import Path


class ExtractionError(Exception):
    """Raised when a document cannot be turned into text."""


_TEXT_SUFFIXES = {".txt", ".vtt", ".md", ".csv"}
_PDF_MIMES = {"application/pdf"}
_DOCX_MIMES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
}


def extract_text(content: bytes, *, filename: str, mime: str | None = None) -> str:
    if not content:
        raise ExtractionError("could not read file")

    suffix = Path(filename).suffix.lower()
    mime = (mime or "").split(";")[0].strip().lower()

    if suffix in _TEXT_SUFFIXES or mime.startswith("text/"):
        text = _decode_text(content)
    elif suffix == ".pdf" or mime in _PDF_MIMES:
        text = _extract_pdf(content)
    elif suffix == ".docx" or mime in _DOCX_MIMES:
        text = _extract_docx(content)
    else:
        raise ExtractionError(f"unsupported file type: {filename}")

    text = text.strip()
    if not text:
        raise ExtractionError("could not read file")
    return text


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader
    from pypdf.errors import FileNotDecryptedError, PdfReadError

    try:
        reader = PdfReader(io.BytesIO(content))
        if getattr(reader, "is_encrypted", False):
            raise ExtractionError("could not read file")
        pages = [(page.extract_text() or "") for page in reader.pages]
    except (FileNotDecryptedError, PdfReadError) as exc:
        raise ExtractionError("could not read file") from exc
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    from docx import Document
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = Document(io.BytesIO(content))
    except (PackageNotFoundError, ValueError) as exc:
        raise ExtractionError("could not read file") from exc

    parts: list[str] = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)
